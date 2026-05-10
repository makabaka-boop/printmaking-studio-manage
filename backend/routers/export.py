from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime
import re
from urllib.parse import quote
from backend.database.database import get_db
from backend.models import models

router = APIRouter(prefix="/api/export", tags=["导出"])

@router.get("/artwork/{artwork_id}")
def export_artwork_ledger(artwork_id: int, db: Session = Depends(get_db)):
    artwork = db.query(models.Artwork).filter(models.Artwork.id == artwork_id).first()
    if not artwork:
        raise HTTPException(status_code=404, detail="作品不存在")
    
    def sanitize_filename(name):
        if not name:
            return "unknown"
        return re.sub(r'[\\/:*?"<>|]', '_', name)
    
    wb = Workbook()
    ws = wb.active
    sheet_title = sanitize_filename(artwork.name)[:28] + " 版次台账" if len(artwork.name or "") > 0 else "版次台账"
    ws.title = sheet_title[:31]
    
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center")
    
    ws.merge_cells('A1:I1')
    ws['A1'] = f"《{artwork.name}》 Edition 台账"
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = center_alignment
    
    info_start_row = 3
    ws[f'A{info_start_row}'] = "作品名称"
    ws[f'B{info_start_row}'] = artwork.name or "-"
    ws[f'A{info_start_row + 1}'] = "版种"
    ws[f'B{info_start_row + 1}'] = artwork.print_type or "-"
    ws[f'A{info_start_row + 2}'] = "成品尺寸"
    ws[f'B{info_start_row + 2}'] = artwork.finished_size or "-"
    ws[f'A{info_start_row + 3}'] = "计划版数"
    ws[f'B{info_start_row + 3}'] = artwork.planned_edition
    ws[f'A{info_start_row + 4}'] = "签名规则"
    ws[f'B{info_start_row + 4}'] = artwork.signature_rule or "-"
    
    header_row = info_start_row + 6
    headers = ["序号", "印制日期", "使用版次", "试印张数", "正印张数", "废张数", "纸张", "用量", "备注"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
    
    batches = db.query(models.PrintBatch).filter(models.PrintBatch.artwork_id == artwork_id).order_by(models.PrintBatch.print_date).all()
    
    total_test = 0
    total_good = 0
    total_waste = 0
    
    for idx, batch in enumerate(batches, 1):
        row = header_row + idx
        try:
            plates_str = ", ".join([f"第{bp.plate.color_number}色" for bp in batch.plates_used])
        except:
            plates_str = "-"
        paper_name = batch.paper.name if batch.paper else "-"
        
        ws.cell(row=row, column=1, value=idx).alignment = center_alignment
        ws.cell(row=row, column=2, value=batch.print_date.strftime("%Y-%m-%d")).alignment = center_alignment
        ws.cell(row=row, column=3, value=plates_str).alignment = center_alignment
        ws.cell(row=row, column=4, value=batch.test_prints).alignment = center_alignment
        ws.cell(row=row, column=5, value=batch.good_prints).alignment = center_alignment
        ws.cell(row=row, column=6, value=batch.waste_prints).alignment = center_alignment
        ws.cell(row=row, column=7, value=paper_name).alignment = center_alignment
        ws.cell(row=row, column=8, value=batch.paper_used).alignment = center_alignment
        ws.cell(row=row, column=9, value=batch.notes or "")
        
        total_test += batch.test_prints
        total_good += batch.good_prints
        total_waste += batch.waste_prints
    
    summary_row = header_row + len(batches) + 2
    ws[f'A{summary_row}'] = "合计"
    ws[f'A{summary_row}'].font = Font(bold=True)
    ws[f'D{summary_row}'] = total_test
    ws[f'E{summary_row}'] = total_good
    ws[f'F{summary_row}'] = total_waste
    
    remaining_row = summary_row + 1
    ws[f'A{remaining_row}'] = "剩余可售版数"
    ws[f'A{remaining_row}'].font = Font(bold=True)
    ws[f'B{remaining_row}'] = max(0, artwork.planned_edition - total_good)
    ws[f'B{remaining_row}'].font = Font(bold=True, color="00B050")
    
    for col in range(1, 10):
        ws.column_dimensions[chr(64 + col)].width = 15
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    safe_name = sanitize_filename(artwork.name or "作品")
    filename = f"{safe_name}_版次台账_{datetime.now().strftime('%Y%m%d')}.xlsx"
    encoded_filename = quote(filename)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )
